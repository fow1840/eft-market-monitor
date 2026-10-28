"""Bounded Excel import. No formula evaluation, network access, or persistence.

The investment ledger's editable bucket targets take precedence over formula
columns, which may have stale Excel caches. Private holdings never become test
fixtures or application logs.
"""
from __future__ import annotations
from datetime import date, datetime
from zoneinfo import ZoneInfo
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import iterparse, ParseError
import openpyxl

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 64
MAX_SHEETS = 20


def _text(value):
    return unicodedata.normalize('NFKC',str(value)).strip() if value is not None else ''


def _header(value):
    return re.sub(r'\s+','',_text(value)).replace('（','(').replace('）',')')


def _number(value, *, percentage=False):
    if isinstance(value,bool) or value is None:
        raise ValueError('金额或比例缺失/格式错误')
    raw=_text(value).replace(',','').replace('，','').replace('￥','').replace('¥','')
    has_pct=raw.endswith('%')
    if has_pct:raw=raw[:-1]
    try:number=Decimal(raw)
    except InvalidOperation as exc:raise ValueError('金额或比例必须为有效数值') from exc
    if not number.is_finite() or number<0:raise ValueError('金额或比例不可为负或非有限值')
    if has_pct:
        if not percentage:raise ValueError('金额不可使用百分号')
        number/=100
    if percentage and number>1:raise ValueError('目标比例使用0—1或百分号，不能猜测百分比单位')
    if not percentage:
        if number>Decimal('1000000000000000'):
            raise ValueError('金额超出支持范围')
        if number != number.quantize(Decimal('.01')):
            raise ValueError('金额最多保留两位小数')
    return number


def _input_bytes(source):
    if isinstance(source,(str,Path)):
        path=Path(source)
        if path.suffix.lower()!='.xlsx':raise ValueError('仅支持.xlsx格式')
        if path.stat().st_size>MAX_UPLOAD_BYTES:raise ValueError('上传文件过大')
        payload=path.read_bytes()
    elif isinstance(source,(bytes,bytearray)):
        payload=bytes(source)
    elif hasattr(source,'read'):
        source.seek(0)
        payload=source.read(MAX_UPLOAD_BYTES+1)
    else:raise ValueError('需要Excel文件、字节或文件流')
    if len(payload)>MAX_UPLOAD_BYTES:raise ValueError('上传文件过大')
    try:
        with ZipFile(BytesIO(payload)) as archive:
            members=archive.infolist()
            if sum(m.file_size for m in members)>MAX_EXPANDED_BYTES or len(members)>2000:
                raise ValueError('Excel展开大小超限')
            if any(m.flag_bits&1 for m in members):raise ValueError('不支持加密Excel')
            # Explicit XML row numbers detect sparse out-of-bound data, which
            # iter_rows(max_row=...) alone would silently ignore.
            for member in members:
                if member.filename.startswith('xl/worksheets/') and member.filename.endswith('.xml'):
                    with archive.open(member) as stream:
                        for event,element in iterparse(stream,events=('start','end')):
                            if event=='start' and element.tag.endswith('}row'):
                                try:row_number=int(element.attrib.get('r','0'))
                                except ValueError as exc:raise ValueError('无效Excel行号') from exc
                                if row_number>MAX_ROWS:raise ValueError('Excel行数超限')
                            if event=='end':element.clear()
    except (BadZipFile, ParseError) as exc:raise ValueError('文件不是有效xlsx') from exc
    return payload


def _rows(sheet):
    # Some spreadsheets omit worksheet dimensions; explicit bounded iteration
    # handles them without trusting max_row/max_column metadata.
    rows=[]
    for row in sheet.iter_rows(min_row=1,max_row=MAX_ROWS+1,max_col=MAX_COLUMNS+1,values_only=True):
        if len(row)>MAX_COLUMNS and any(v is not None for v in row[MAX_COLUMNS:]):
            raise ValueError('Excel列数超限')
        rows.append(tuple(row[:MAX_COLUMNS]))
    if len(rows)>MAX_ROWS and any(v is not None for v in rows[MAX_ROWS]):raise ValueError('Excel行数超限')
    while rows and all(v is None for v in rows[-1]):rows.pop()
    return rows


def _find_header(rows, required):
    for i,row in enumerate(rows[:20]):
        headers={_header(value):j for j,value in enumerate(row) if _text(value)}
        if all(any(alias in headers for alias in group) for group in required):return i,headers
    raise ValueError('未找到所需Excel表头')


def _column(headers, *aliases):
    return next((headers[_header(a)] for a in aliases if _header(a) in headers),None)


def _cell(row, column):return row[column] if column is not None and column<len(row) else None


def _literal(row, cached_row, column, warnings, label):
    value=_cell(row,column)
    if isinstance(value,str) and value.startswith('='):
        value=_cell(cached_row,column)
        if value is None:raise ValueError(f'{label}公式没有缓存值，请在Excel重算或填写数值')
        warnings.append(f'{label}使用Excel公式缓存；请确认已重算后保存')
    return value


def _restriction(row,headers,today):
    period=_text(_cell(row,_column(headers,'持有期','持有期限','最短持有期')))
    restriction=_text(_cell(row,_column(headers,'赎回限制','赎回限制说明')))
    redeem=_text(_cell(row,_column(headers,'可赎回','赎回状态','是否可赎回')))
    earliest_raw=_cell(row,_column(headers,'最早赎回日','最早赎回日期','可赎回日期'))
    earliest=None
    if earliest_raw is not None and _text(earliest_raw):
        try:
            if isinstance(earliest_raw,datetime):earliest=earliest_raw.date()
            elif isinstance(earliest_raw,date):earliest=earliest_raw
            else:earliest=date.fromisoformat(_text(earliest_raw)[:10].replace('/','-'))
        except (ValueError,TypeError):return 'unknown','赎回日期格式无效，需人工确认',period,None
    deny={'否','不可','不可赎回','暂停赎回','受限','锁定','封闭','0','false'}
    allow={'是','可','可赎回','正常','开放','开放赎回','1','true','无限制','无','无赎回限制'}
    restriction_lower=restriction.lower();redeem_lower=redeem.lower()
    if earliest and earliest>today:
        return 'locked','最早赎回日期尚未到达',period,earliest.isoformat()
    if redeem_lower in deny or restriction_lower in deny or any(word in restriction+' '+period for word in ('不可赎回','暂停赎回','锁定中','封闭期内','禁止赎回','未满持有期','未到期','受限')):
        return 'locked','表内标记赎回受限',period,earliest.isoformat() if earliest else None
    # Conflicting explicit statements are never interpreted as permission.
    if redeem_lower in allow and restriction_lower not in allow|{''}:
        return 'unknown','可赎回与限制说明存在不确定性，需人工确认',period,earliest.isoformat() if earliest else None
    if redeem_lower in allow or restriction_lower in allow:
        return 'free','表内明确标记可赎回；执行前仍需核对费用',period,earliest.isoformat() if earliest else None
    return 'unknown','持有期/赎回限制信息不足，需人工确认',period,earliest.isoformat() if earliest else None


def parse_holdings_excel(source, *, today=None):
    """Return holdings, bucket targets, warnings and non-personal metadata.

    Money and fractions use Decimal. Holding statuses are free/locked/unknown;
    missing restrictions are never inferred from a fund's name or asset bucket.
    """
    today=today or datetime.now(ZoneInfo('Asia/Shanghai')).date()
    payload=_input_bytes(source)
    try:
        book=openpyxl.load_workbook(BytesIO(payload),read_only=True,data_only=False,keep_links=False)
        cached=openpyxl.load_workbook(BytesIO(payload),read_only=True,data_only=True,keep_links=False)
    except Exception as exc:raise ValueError('无法读取Excel工作簿') from exc
    try:
        if len(book.sheetnames)>MAX_SHEETS:raise ValueError('Excel工作表数量超限')
        if not {'当前持仓','目标组合'}<=set(book.sheetnames):raise ValueError('需要“当前持仓”和“目标组合”工作表')
        holdings_rows=_rows(book['当前持仓']);cached_holdings=_rows(cached['当前持仓'])
        target_rows=_rows(book['目标组合']);cached_targets=_rows(cached['目标组合'])
        index,headers=_find_header(holdings_rows,[('基金名称','基金名'),('当前金额(元)','金额','金额(元)','当前金额'),('资产桶','资产类别')])
        name_col=_column(headers,'基金名称','基金名');amount_col=_column(headers,'当前金额(元)','金额','金额(元)','当前金额');bucket_col=_column(headers,'资产桶','资产类别')
        account_col=_column(headers,'账户','账号');themes_col=_column(headers,'重复/重叠组','重复重叠组','主题','重叠组')
        fund_target_col=_column(headers,'基金目标仓位','基金目标比例','目标仓位','目标占比')
        warnings=[];holdings=[]
        for position,row in enumerate(holdings_rows[index+1:],start=index+1):
            name=_text(_cell(row,name_col));amount=_cell(row,amount_col);bucket=_text(_cell(row,bucket_col))
            if name in {'合计','总计','总额'}:continue
            if not name and amount is None and not bucket:continue
            if not name or not bucket:raise ValueError(f'当前持仓第{position+1}行缺少基金名称或资产桶')
            cached_row=cached_holdings[position] if position<len(cached_holdings) else ()
            value=_number(_literal(row,cached_row,amount_col,warnings,f'第{position+1}行金额'))
            status,reason,period,earliest=_restriction(row,headers,today)
            themes=[part.strip() for part in re.split('[;；\n]+',_text(_cell(row,themes_col))) if part.strip()]
            entry=dict(id=f'当前持仓:{position+1}',name=name,account=_text(_cell(row,account_col)),amount=value,bucket=bucket,themes=themes,
                       holding_status=status,restriction_reason=reason,holding_period=period,earliest_redemption=earliest,
                       subcategory=_text(_cell(row,_column(headers,'细分类'))),role=_text(_cell(row,_column(headers,'核心/卫星'))))
            if fund_target_col is not None and _cell(row,fund_target_col) is not None:
                entry['fund_target']=_number(_literal(row,cached_row,fund_target_col,warnings,'基金目标比例'),percentage=True)
            holdings.append(entry)
        if not holdings:raise ValueError('没有可分析的持仓')
        target_index,target_headers=_find_header(target_rows,[('资产桶','资产类别'),('目标比例(可编辑)','目标仓位','目标占比','目标比例')])
        target_bucket_col=_column(target_headers,'资产桶','资产类别')
        target_col=_column(target_headers,'目标比例(可编辑)','目标仓位','目标占比','目标比例')
        targets={}
        for position,row in enumerate(target_rows[target_index+1:],start=target_index+1):
            bucket=_text(_cell(row,target_bucket_col))
            if not bucket or bucket in {'合计','总计','总额'}:continue
            if bucket in targets:raise ValueError('目标组合资产桶重复')
            cached_row=cached_targets[position] if position<len(cached_targets) else ()
            targets[bucket]=_number(_literal(row,cached_row,target_col,warnings,'目标比例'),percentage=True)
        if not targets or sum(targets.values())!=Decimal('1'):raise ValueError('目标比例合计必须为100%')
        if any(entry['bucket'] not in targets for entry in holdings):raise ValueError('当前持仓资产桶不在目标组合中，请补充目标（允许0%）')
        if sum(entry['amount'] for entry in holdings)<=0:raise ValueError('持仓总金额必须大于零')
        if any(entry['holding_status']=='unknown' for entry in holdings):warnings.append('部分持有期/赎回限制缺失或不明确，需人工确认；基金名称不能证明可赎回')
        normalized=[re.sub(r'\s+','',entry['name']) for entry in holdings]
        if len(set(normalized))<len(normalized):warnings.append('存在同名基金跨账户或重复行，请核对；导入保留各行，不擅自删除或重复汇总')
        return dict(holdings=holdings,targets=targets,warnings=list(dict.fromkeys(warnings)),metadata={'holdings_sheet':'当前持仓','targets_sheet':'目标组合','holding_rows':len(holdings),'target_buckets':len(targets),'formulas_evaluated':False})
    finally:
        book.close();cached.close()
